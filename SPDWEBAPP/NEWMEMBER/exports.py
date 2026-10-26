"""Shared approved-marks workbook with identity-safe blocks and literal spreadsheet text."""
from collections import defaultdict, Counter
from io import BytesIO
from django.core.exceptions import PermissionDenied, ValidationError
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter
from AUTHENTICATE.models import Brother_Profile, PClass
from .models import NewMember_Mark_Event_and_Request as Mark
from .operations import require_marks, name_for


def class_choices():
    present = set(Brother_Profile.objects.exclude(pclass__isnull=True).values_list('pclass', flat=True))
    return [value for value, _ in reversed(PClass.choices) if value in present]


def text(cell, value):
    cell.value, cell.data_type = str(value), 's'


def export_marks(user, pclass):
    _, roles = require_marks(user)
    if 'NEW_MEMBER' in roles and pclass != 'Current':
        raise PermissionDenied('You do not have permission to export marks for previous classes.')
    if pclass is not None and pclass not in ['Current', *class_choices()]:
        raise ValidationError({'pclass': 'Select an available class.'})
    profiles = Brother_Profile.objects.filter(roles__name='NEW_MEMBER') if pclass == 'Current' else Brother_Profile.objects.filter(pclass=pclass)
    ids = profiles.values_list('user_id', flat=True)
    rows = Mark.objects.filter(target_user_id__in=ids, mark_approval_status='approved').select_related('target_user', 'requesting_user').order_by('mark_event_date', 'pk')
    grouped = defaultdict(list)
    for mark in rows:
        grouped[mark.target_user_id].append(mark)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'New Member Marks'
    palette = ['FFC7CE', 'C6EFCE', 'FFEB9C', 'C9DAF8']
    names = Counter(name_for(marks[0].target_user) for marks in grouped.values())
    for index, (target_id, marks) in enumerate(sorted(grouped.items(), key=lambda item: (name_for(item[1][0].target_user), item[0]))):
        first = index * 5 + 1
        fill = PatternFill('solid', fgColor=palette[index % len(palette)])
        for row, label in [(1, name_for(marks[0].target_user) + (f' (member #{target_id})' if names[name_for(marks[0].target_user)] > 1 else '')), (2, f'Total: {sum(mark.mark_value for mark in marks)}')]:
            sheet.merge_cells(start_row=row, start_column=first, end_row=row, end_column=first+3)
            cell = sheet.cell(row, first)
            text(cell, label)
            cell.fill, cell.alignment = fill, Alignment(horizontal='center')
        for offset, label in enumerate(['Date', 'Requesting Member', 'Reason', 'Mark Value']):
            cell = sheet.cell(3, first+offset)
            text(cell, label)
            cell.fill = fill
            sheet.column_dimensions[get_column_letter(first+offset)].width = [12, 20, 30, 10][offset]
        for row, mark in enumerate(marks, 4):
            for offset, value in enumerate([mark.mark_event_date.isoformat(), name_for(mark.requesting_user), mark.mark_event_title]):
                text(sheet.cell(row, first+offset), value)
            sheet.cell(row, first+3, mark.mark_value)
        sheet.column_dimensions[get_column_letter(first+4)].width = 2
    buffer = BytesIO()
    workbook.save(buffer)
    name = f'Approved_{pclass}_Class_Newmember_Marks.xlsx' if pclass else 'Approved_Newmember_Marks.xlsx'
    response = HttpResponse(buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{name}"'
    response['Cache-Control'] = 'no-store'
    return response
