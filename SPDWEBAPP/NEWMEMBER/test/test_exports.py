from io import BytesIO

from openpyxl import load_workbook
from django.urls import reverse

from .support import MarkCase


class ExportTests(MarkCase):
    def setUp(self):
        super().setUp()
        self.a = self.user('a', ('ACTIVE',), profile_fields={'firstName': 'Alex', 'lastName': 'Able'})
        self.t1 = self.user('t1', ('NEW_MEMBER',), profile_fields={'firstName': 'Tess', 'lastName': 'Able'})
        self.t2 = self.user('t2', ('NEW_MEMBER',), profile_fields={'firstName': 'Zoe', 'lastName': 'Zed'})
        self.b = self.user('b', ('ALUMNI',), profile_fields={'pclass': 'Alpha'})

    def workbook(self, actor):
        self.login(actor)
        before = self.capture()
        response = self.client.get(reverse('export_newmember_marks'), {'pclass': 'Current'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertEqual(response['Content-Disposition'], 'attachment; filename="Approved_Current_Class_Newmember_Marks.xlsx"')
        workbook = load_workbook(BytesIO(response.content), data_only=False)
        self.assertEqual(workbook.sheetnames, ['New Member Marks'])
        self.unchanged(before)
        self.identity(actor)
        return workbook.active

    def test_export_current_workbook_has_signed_approved_values(self):
        self.mark(self.a, self.t1, status='approved', value=5, day='2026-10-14', title='Early')
        self.mark(self.a, self.t1, status='approved', value=-8, title='Late')
        self.mark(self.a, self.t2, status='approved', value=2, title='Second member')
        self.mark(self.a, self.t1, status='requested', value=99, title='Pending noise')
        self.mark(self.a, self.t1, status='denied', value=99, title='Denied noise')
        self.mark(self.a, self.b, status='approved', value=90, title='Historical noise')
        sheet = self.workbook(self.a)
        self.assertEqual(sheet['A1'].value, 'Tess Able')
        self.assertEqual(sheet['F1'].value, 'Zoe Zed')
        self.assertEqual(sheet['A2'].value, 'Total: -3')
        self.assertEqual(sheet['F2'].value, 'Total: 2')
        headers = ['Date', 'Requesting Member', 'Reason', 'Mark Value']
        self.assertEqual([sheet.cell(3, col).value for col in range(1, 5)], headers)
        self.assertEqual([sheet.cell(3, col).value for col in range(6, 10)], headers)
        self.assertEqual([sheet.cell(4, col).value for col in range(1, 5)],
                         ['2026-10-14', 'Alex Able', 'Early', 5])
        self.assertEqual([sheet.cell(5, col).value for col in range(1, 5)],
                         ['2026-10-16', 'Alex Able', 'Late', -8])
        self.assertEqual([sheet.cell(4, col).value for col in range(6, 10)],
                         ['2026-10-16', 'Alex Able', 'Second member', 2])
        for address in ['D4', 'D5', 'I4']:
            self.assertEqual(sheet[address].data_type, 'n')
        for address in ['A1', 'A2', 'A4', 'B4', 'C4']:
            self.assertEqual(sheet[address].data_type, 's')
        self.assertTrue(all(sheet.cell(row, 5).value is None for row in range(1, 6)))
        values = [c.value for row in sheet for c in row]
        for noise in ['Historical noise', 'Denied noise', 'Pending noise']:
            self.assertNotIn(noise, values)

    def test_new_member_current_export_is_allowed(self):
        for index, roles in enumerate([('NEW_MEMBER',), ('NEW_MEMBER', 'NM_BOARD')]):
            actor = self.user('viewer-' + str(index), roles,
                              profile_fields={'firstName': 'Viewer' + str(index), 'lastName': 'Member'})
            mine = self.mark(self.a, actor, status='approved', value=-3, title='Own ' + str(index))
            other = self.mark(self.a, self.t1, status='approved', value=2, title='Other ' + str(index))
            historical = self.mark(self.a, self.b, status='approved', value=99, title='Historical ' + str(index))
            self.mark(self.a, actor, status='requested', value=50, title='Pending ' + str(index))
            with self.subTest(roles=roles):
                sheet = self.workbook(actor)
                values = [cell.value for row in sheet for cell in row]
                self.assertIn('Own ' + str(index), values)
                self.assertIn('Other ' + str(index), values)
                self.assertIn('Total: -3', values)
                self.assertNotIn('Historical ' + str(index), values)
                self.assertNotIn('Pending ' + str(index), values)
                for row in sheet:
                    for cell in row:
                        if cell.value in [-3, 2] and isinstance(cell.value, int):
                            self.assertEqual(cell.data_type, 'n')

    def test_empty_current_export_is_a_valid_workbook(self):
        self.mark(self.a, self.t1, status='requested')
        self.mark(self.a, self.t2, status='denied')
        self.mark(self.a, self.b, status='approved', title='Historical')
        sheet = self.workbook(self.a)
        self.assertEqual((sheet.max_row, sheet.max_column), (1, 1))
        self.assertTrue(all(c.value is None for row in sheet for c in row))
        self.t1.brother_profile.roles.clear()
        self.t2.brother_profile.roles.clear()
        sheet = self.workbook(self.a)
        self.assertEqual((sheet.max_row, sheet.max_column), (1, 1))
        self.assertTrue(all(c.value is None for row in sheet for c in row))

    def test_export_requester_without_profile_uses_username(self):
        requester = self.user('requester-fallback', with_profile=False, username='requester-fallback')
        mark = self.mark(requester, self.t1, status='approved', value=-3, title='Fallback reason')
        self.mark(self.a, self.b, status='approved', title='Unrelated historical')
        sheet = self.workbook(self.a)
        self.assertEqual([sheet.cell(4, c).value for c in range(1, 5)],
                         ['2026-10-16', requester.username, 'Fallback reason', -3])
        self.assertEqual(sheet['A1'].value, 'Tess Able')
        self.assertEqual(sheet['A2'].value, 'Total: -3')
        self.assertEqual(sheet['B4'].data_type, 's')
        self.assertEqual(sheet['D4'].data_type, 'n')
