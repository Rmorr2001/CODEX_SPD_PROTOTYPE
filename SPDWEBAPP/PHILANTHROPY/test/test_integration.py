from io import BytesIO

from openpyxl import load_workbook

from SPDWEBAPP_CORE.testing.factories import make_user
from SPDWEBAPP_CORE.testing.helpers import at_time, login_actor
from SPDWEBAPP_CORE.testing.state import assert_unchanged

from .support import (APPROVE, DASHBOARD, EXPORT, SUBMIT, T1, Event,
                      PhilanthropyCase, history, make_event, payload)


class ApprovedProjectionTests(PhilanthropyCase):
    def test_submit_then_approve_updates_all_approved_projections_once(self):
        a = make_user(username='a', roles=('ACTIVE',),
                       profile_fields={'firstName': 'Alex', 'lastName': 'One'})
        b = make_user(username='b', roles=('ACTIVE',),
                       profile_fields={'firstName': 'Blair', 'lastName': 'Two'})
        chair = make_user(username='chair', roles=('ACTIVE', 'PHIL_CHAIR'),
                           profile_fields={'firstName': 'Casey', 'lastName': 'Chair'})
        control = make_event(b, title='Control', hours=7, status='approved')
        login_actor(self.client, a)
        before = self.capture(a, b, control)
        with at_time(T1):
            response = self.client.post(SUBMIT, payload())
        self.assertRedirects(response, DASHBOARD, fetch_redirect_response=False)
        event = Event.objects.exclude(pk=control.pk).get()
        self.assert_requested(event, a)
        dashboard = self.client.get(DASHBOARD, {'year': '2026'})
        self.assertEqual(dashboard.status_code, 200)
        self.assertEqual(list(dashboard.context['user_events']), [event])
        self.assertEqual(dashboard.context['user_total_hours'], 0)
        own = self.client.get(history(a))
        self.assertEqual(own.status_code, 200)
        self.assertEqual(own.context['lifetime_total'], 0)
        self.unchanged(before)
        login_actor(self.client, chair)
        before = self.capture(a, b, chair, event, control)
        wb = load_workbook(BytesIO(self.client.get(EXPORT).content), data_only=False)
        try:
            values = [cell.value for row in wb.active for cell in row]
            self.assertIn('Control', values)
            self.assertNotIn('Park', values)
        finally:
            wb.close()
        with at_time(T1):
            response = self.client.post(APPROVE, {'event_id': event.pk, 'action': 'approve'})
        self.assertEqual(response.status_code, 200)
        event.refresh_from_db()
        self.assertEqual(event.philanthropy_approval_status, 'approved')
        self.assertEqual(event.philanthropy_event_approver_name, 'Casey Chair')
        self.assertEqual(event.philanthropy_event_submission_approval_date, T1)
        # Own projections use A; take the next scope after the legitimate login changes.
        assert_unchanged(before, except_fields={event: {
            'philanthropy_approval_status', 'philanthropy_event_approver_name',
            'philanthropy_event_submission_approval_date'}})
        wb = load_workbook(BytesIO(self.client.get(EXPORT).content), data_only=False)
        try:
            ws = wb.active
            self.assertEqual(ws['A1'].value, 'Alex One')
            self.assertEqual(ws['A2'].value, 'Total: 2')
            self.assertEqual((ws['C4'].value, ws['C4'].data_type), (2, 'n'))
            self.assertEqual([cell.value for row in ws for cell in row].count('Park'), 1)
        finally:
            wb.close()
        login_actor(self.client, a)
        after_decision = self.capture(a, b, chair, event, control)
        dashboard = self.client.get(DASHBOARD, {'year': '2026'})
        own = self.client.get(history(a))
        self.assertEqual(dashboard.context['user_total_hours'], 2)
        self.assertEqual(dashboard.context['total_chapter_hours'], 9)
        self.assertEqual(own.context['lifetime_total'], 2)
        self.assertEqual(own.context['academic_year_totals'][0]['events'], [event])
        self.assertEqual(event.philanthropy_event_submission_request_date, T1)
        self.assertEqual(Event.objects.count(), 2)
        self.unchanged(after_decision)
