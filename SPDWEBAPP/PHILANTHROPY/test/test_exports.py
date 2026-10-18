from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from SPDWEBAPP_CORE.testing.factories import make_user
from SPDWEBAPP_CORE.testing.helpers import login_actor

from .support import EXPORT, Event, PhilanthropyCase, make_event


class ExportTests(PhilanthropyCase):
    def setUp(self):
        super().setUp()
        self.c = make_user(username='chair', roles=('ACTIVE', 'PHIL_CHAIR'))
        self.a = make_user(username='a', roles=('ACTIVE',),
                           profile_fields={'firstName': 'Alex', 'lastName': 'One'})
        self.b = make_user(username='b', roles=('ACTIVE',),
                           profile_fields={'firstName': 'Blair', 'lastName': 'Two'})
        login_actor(self.client, self.c)

    def workbook(self):
        before_files = list(self.media_root.rglob('*'))
        response = self.client.get(EXPORT)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'],
                         'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertEqual(response['Content-Disposition'],
                         'attachment; filename=approved_philanthropy_hours.xlsx')
        wb = load_workbook(BytesIO(response.content), data_only=False)
        self.addCleanup(wb.close)
        self.assertEqual(wb.sheetnames, ['Philanthropy Hours'])
        self.assertEqual(list(self.media_root.rglob('*')), before_files)
        return wb.active

    def test_xlsx_content_totals_and_typed_hours(self):
        e2 = make_event(self.a, title='Earlier service', hours=2,
                        day=date(2025, 7, 1), status='approved')
        e3 = make_event(self.a, title='Later service', hours=3,
                        day=date(2025, 7, 2), status='approved',
                        philanthropy_event_approver_name='Casey Chair')
        f10 = make_event(self.b, title='Blair service', hours=10,
                         day=date(2025, 7, 3), status='approved',
                         philanthropy_event_approver_name='Casey Chair')
        requested = make_event(self.a, title='Pending exclusion', hours=99)
        denied = make_event(self.b, title='Denied exclusion', hours=99, status='denied')
        before = self.capture(self.a, self.b, self.c, e2, e3, f10, requested, denied)
        ws = self.workbook()
        self.assertEqual((ws['A1'].value, ws['F1'].value), ('Alex One', 'Blair Two'))
        self.assertEqual((ws['A2'].value, ws['F2'].value), ('Total: 5', 'Total: 10'))
        self.assertTrue({'A1:D1', 'A2:D2', 'F1:I1', 'F2:I2'} <=
                        {str(r) for r in ws.merged_cells.ranges})
        for start in (1, 6):
            self.assertEqual([ws.cell(3, start + offset).value for offset in range(4)],
                             ['Date', 'Event', 'Hours', 'Approved By'])
        self.assertEqual([ws.cell(4, i).value for i in range(1, 5)],
                         ['2025-07-01', 'Earlier service', 2, '—'])
        self.assertEqual([ws.cell(5, i).value for i in range(1, 5)],
                         ['2025-07-02', 'Later service', 3, 'Casey Chair'])
        self.assertEqual([ws.cell(4, i).value for i in range(6, 10)],
                         ['2025-07-03', 'Blair service', 10, 'Casey Chair'])
        for address, value in (('C4', 2), ('C5', 3), ('H4', 10)):
            self.assertEqual((ws[address].value, ws[address].data_type), (value, 'n'))
        flattened = [cell.value for row in ws for cell in row]
        for title in ('Earlier service', 'Later service', 'Blair service'):
            self.assertEqual(flattened.count(title), 1)
        for title in ('Pending exclusion', 'Denied exclusion'):
            self.assertNotIn(title, flattened)
        self.assertTrue(all(ws.cell(row, 5).value is None for row in range(1, 6)))
        self.unchanged(before)
        self.assertEqual(Event.objects.count(), 5)

    def test_xlsx_total_fill_thresholds_and_event_wrapping(self):
        d = make_user(username='d', roles=('ACTIVE',),
                       profile_fields={'firstName': 'Dana', 'lastName': 'Three'})
        events = [make_event(user, hours=hours, status='approved', day=date(2025, 7, i))
                  for i, (user, hours) in enumerate(((self.a, 4), (self.b, 5), (d, 10)), 1)]
        before = self.capture(self.a, self.b, self.c, d, *events)
        ws = self.workbook()
        for start, total, color in ((1, 4, 'FFC7CE'), (6, 5, 'FFEB9C'), (11, 10, 'C6EFCE')):
            self.assertEqual(ws.cell(2, start).value, f'Total: {total}')
            self.assertEqual(ws.cell(1, start).fill.fgColor.rgb[-6:], color)
            self.assertEqual(ws.cell(2, start).fill.fgColor.rgb[-6:], color)
            self.assertIs(ws.cell(4, start + 1).alignment.wrap_text, True)
            self.assertIsNone(ws.cell(4, start + 4).value)
        self.unchanged(before)
        self.assertEqual(Event.objects.count(), 3)

    def test_unicode_and_plain_multiline_export_text_round_trips(self):
        self.a.brother_profile.firstName = 'Δ'
        self.a.brother_profile.lastName = 'Member'
        self.a.brother_profile.save()
        event = make_event(self.a, title='Park\ncleanup\tteam', hours=2, status='approved',
                           philanthropy_event_approver_name='Zoë Chair')
        before = self.capture(self.a, self.b, self.c, event)
        ws = self.workbook()
        self.assertEqual(ws['A1'].value, 'Δ Member')
        self.assertEqual(ws['B4'].value, 'Park\ncleanup\tteam')
        self.assertEqual(ws['D4'].value, 'Zoë Chair')
        for address in ('A1', 'B4', 'D4'):
            self.assertEqual(ws[address].data_type, 's')
        self.assertEqual((ws['C4'].value, ws['C4'].data_type), (2, 'n'))
        self.unchanged(before)

    def test_export_with_no_approved_events_is_readable_empty_workbook(self):
        requested = make_event(self.a, title='Pending')
        denied = make_event(self.b, title='Denied', status='denied')
        before = self.capture(self.a, self.b, self.c, requested, denied)
        ws = self.workbook()
        self.assertTrue(all(cell.value is None for row in ws for cell in row))
        self.assertEqual(list(ws.merged_cells.ranges), [])
        self.unchanged(before)
        self.assertEqual(Event.objects.count(), 2)
